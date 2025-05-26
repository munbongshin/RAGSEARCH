
import sys, os
#sys.path.append('C:/Dev/vueprj2')  # 프로젝트 루트 경로 추가
sys.path.append(os.path.join(os.path.dirname(__file__), 'backend/app'))
from backend.app.ExtractTextFromFile import ExtractTextFromFile
from backend.app.CustomSentenceTransformerEmbeddings import CustomSentenceTransformerEmbeddings as CSTFM
from backend.app.QuoteExtractor import QuoteExtractor as QuoExt

from dotenv import load_dotenv
from pathlib import Path
import chromadb
import uuid
import logging, traceback, re
from openpyxl import load_workbook
from pptx import Presentation
import json
#from langchain_community.vectorstores import Chroma
from langchain_chroma import Chroma
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.docstore.document import Document
import warnings
from konlpy.tag import Kkma
from sklearn.preprocessing import MinMaxScaler
import numpy as np
from langchain.chains.summarize import load_summarize_chain
from langchain.prompts import PromptTemplate
from langchain.text_splitter import TokenTextSplitter
from typing import Union, List, Any
# groq 라이브러리 import 수정
import groq
from logging.handlers import RotatingFileHandler

      
logging.basicConfig(level=logging.DEBUG)
logging.getLogger("langchain").setLevel(logging.ERROR)
logger = logging.getLogger(__name__)

class ChromaDbManager:
    distance_metric = "l2"  # 클래스 속성으로 이동
    def __init__(self, persist_directory=""):
        try:
            project_root = Path(__file__).parent
            env_path = project_root / '.env'
            load_dotenv(dotenv_path=env_path)
            current_dir = os.path.dirname(os.path.abspath(__file__))
            MODEL_PATH = os.path.join(os.path.dirname(current_dir), "chroma_db")
            self.persist_directory = MODEL_PATH if not persist_directory else persist_directory
            self.config_file = "chroma_config.json"
            self.embeddings = CSTFM()
            self.client = self._create_client()
            self.vectordb = None
            #self.docnum = os.environ.get("DOC_NUM")
            #self.chunk_size = os.environ.get("CHUNK_SIZE")
            #self.chunk_overlap = os.environ.get("CHUNK_OVERLAP")
            self.docnum = int(os.getenv("DOC_NUM", "3"))
            self.chunk_size = int(os.getenv("CHUNK_SIZE", "1000"))
            self.chunk_overlap = int(os.getenv("CHUNK_OVERLAP", "200"))
            self.extractor = ExtractTextFromFile()
        except Exception as e:
            error_message = f"ChromaDbManager 초기화 오류: {e}"
            logger.error(error_message)
            raise

    def _create_client(self):
        return chromadb.PersistentClient(path=self.persist_directory)
    
    def set_return_docnum(self, docnum):
        self.docnum = docnum
       
    def load_config(self):
        try:
            if os.path.exists(self.config_file):
                with open(self.config_file, 'r') as f:
                    return json.load(f)
            return {}
        except Exception as e:
            error_message = f"CharomaDB load_config 오류 발생: {e}"
            logger.debug(f"{error_message}")
            return []
        
    def get_persist_directory(self):
        return self.persist_directory

    def save_config(self):
        try:
            config = {'persist_directory': self.persist_directory}
            with open(self.config_file, 'w') as f:
                json.dump(config, f)
        except Exception as e:
            error_message = f"charomadb save_config 오류 발생: {e}"
            logger.debug(f"{error_message}")
            return []
    
        
    def get_or_create_collection(self, collection_name):
            try:
                if self.vectordb is None or self.vectordb._collection.name != collection_name:
                    self.vectordb = Chroma(
                        collection_name=collection_name,
                        embedding_function=self.embeddings,
                        persist_directory=self.persist_directory,
                    )
                return self.vectordb
            except Exception as e:
                error_message = f"get_or_create_collection 오류: {e}"
                logger.error(error_message)
                return None   
    
    def extract_text_from_file(self, file, file_name):
        return self.extractor.extract_text_from_file(file, file_name)
    
    def set_persist_directory(self, new_directory):
        try:
            if self.persist_directory != new_directory:
                self.persist_directory = new_directory
                # 기존 클라이언트를 닫고 새로운 클라이언트를 생성
                if hasattr(self, 'client'):
                    self.client.close()
                self.client = self._create_client()
        except Exception as e:
            error_message = f"set_persist_directory 오류 발생: {e}"
            logger.debug(f"{error_message}")
                        
    def create_collection(self, collection_name):       
        try:
            
            self.client.create_collection(name=collection_name)
            #logger.debug(f"create collection info: {self.client.get_collection(name=collection_name).count()}")
            return True
        except ValueError as e:
            return False

    def delete_collection(self, collection_name):
        try:
            self.client.delete_collection(name=collection_name)
            #logger.debug(f"delete collection info: {self.client.get_collection(name=collection_name).count()}")
            return f"Collection '{collection_name}' deleted successfully."
        except ValueError as e:
            return f"Error deleting collection: {e}"

    def list_collections(self):
        try:
            collections = self.client.list_collections()
            return [collection.name for collection in collections]
        except Exception as e:
            error_message = f"list_collection 오류 발생: {e}"
            logger.debug(f"{error_message}")
            return []
        
    def get_collection_schema(self, collection_name):
        """
        ChromaDB collection의 필드 구조를 조회하는 함수
        
        Args:
            collection_name: 조회할 collection 이름
        Returns:
            collection의 필드 구조 정보
        """
        try:
            collection = self.client.get_collection(name=collection_name)
            if not collection:
                print(f"컬렉션을 찾을 수 없습니다: {collection_name}")
                return None

            # collection의 데이터 조회
            results = collection.get()
            
            def get_field_info(data):
                """중첩된 필드의 타입과 구조를 분석"""
                if isinstance(data, dict):
                    return {k: get_field_info(v) for k, v in data.items()}
                elif isinstance(data, list) and data:
                    return {
                        'type': 'list',
                        'element_type': type(data[0]).__name__,
                        'sample_structure': get_field_info(data[0]) if isinstance(data[0], (dict, list)) else None
                    }
                else:
                    return {'type': type(data).__name__}

            # 결과가 딕셔너리인 경우
            if isinstance(results, dict):
                field_structure = {}
                for field, value in results.items():
                    field_structure[field] = get_field_info(value)
            else:
                # 결과가 리스트인 경우
                field_structure = {'documents': get_field_info(results)}
                
            schema_info = {
                "collection_name": collection_name,
                "field_structure": field_structure
            }
            
            return schema_info
            
        except Exception as e:
            print(f"오류 발생: {str(e)}")
            return None

    def print_schema_info(self, schema_info):
        """
        필드 구조 정보를 출력하는 함수
        """
        if not schema_info:
            print("스키마 정보를 사용할 수 없습니다")
            return
        
        def print_structure(structure, indent=0):
            """재귀적으로 필드 구조를 출력"""
            prefix = "  " * indent
            if isinstance(structure, dict):
                if 'type' in structure:
                    if structure['type'] == 'list':
                        print(f"{prefix}타입: {structure['type']}")
                        print(f"{prefix}요소 타입: {structure['element_type']}")
                        if structure['sample_structure']:
                            print(f"{prefix}요소 구조:")
                            print_structure(structure['sample_structure'], indent + 1)
                    else:
                        print(f"{prefix}타입: {structure['type']}")
                else:
                    for key, value in structure.items():
                        print(f"{prefix}{key}:")
                        print_structure(value, indent + 1)
                        
        print("\n=== 컬렉션 스키마 정보 ===")
        print(f"컬렉션 이름: {schema_info['collection_name']}")
        print("\n필드 구조:")
        print_structure(schema_info['field_structure']) 
        
    
            
    def view_collection_content(self, collection_name, docnum=100):
        try:
            collection = self.client.get_collection(name=collection_name)
            
            # 컬렉션의 총 문서수 확인
            total_count = collection.count()
            #logger.debug(f"Total items in collection: {total_count}")

            # 한 번에 모든 데이터 가져오기, 서버 측에서 limit 적용
            items = collection.get(limit=docnum)
            
            if not items['ids']:
                return f"Collection '{collection_name}' is empty or data could not be retrieved."

            content = []
            for i in range(len(items['ids'])):
                item_info = {
                    "ID": items['ids'][i],
                    "Metadata": items['metadatas'][i] if items['metadatas'] else "No metadata",
                    "Document": items['documents'][i] if items['documents'] else "No document"
                }
                content.append(item_info)

            return content

        except ValueError as e:
            return f"Error viewing collection: {e}"
        except Exception as e:
            return f"Unexpected error occurred: {e}"

    


    def get_subdirectories(path):
        try:
            return [d for d in os.listdir(path) if os.path.isdir(os.path.join(path, d))]
        except PermissionError:
            print(f"Permission Denied", f"Permission denied to access {path}")
        return []
    
   
   
    def verify_storage(self, collection_name):
        try:
            collection = self.client.get_collection(name=collection_name)
            count = collection.count()
            #logger.debug(f"Verification: Collection '{collection_name}' has {count} items.")
            if count > 0:
                sample = collection.peek()
                #logger.debug(f"Sample item: {sample}")
            return count
        except Exception as e:
            error_message = f"verify_storage 오류 발생: {e}"
            logger.debug(f"{error_message}")
            return []
            
    def test_embedding():
        try:
            #embeddings = FastEmbedEmbeddings(model_name="BAAI/bge-large-en-v1.5")
            #embeddings = OllamaEmbeddings(model="llama3:instruct")
            #embeddings = OllamaEmbeddings(model="nomic-embed-text:latest")
            embeddings = CSTFM()
            test_text = "This is a test sentence."
            #logger.debug(test_text)
            result = embeddings.embed_query(test_text)
            #logger.debug(f"Embedding dimension: {len(result)}")
            #logger.debug(f"First few values: {result[:5]}")
        except Exception as e:
            error_message = f"test_embedding 오류 발생: {e}"
            logger.debug(f"{error_message}")
            return []
    
    def store_in_chroma(self, text, filename, collection_name, progress_callback=None):
        try:
            collection = self.client.get_or_create_collection(collection_name)
            
            existing_docs = self.get_ids_by_source(collection_name=collection_name, source=filename)
            if existing_docs[0]:
                messagebox.showinfo("Success", f"문서 '{filename}'이(가) 이미 존재합니다. 업데이트합니다.")
                chunks = self.chunk_text(text,self.chunk_size)
                total_chunks = len(chunks)                    
                for i, chunk in enumerate(chunks):
                    collection.update(
                        documents=[chunk],
                        metadatas=[{"source": filename}],
                        ids=existing_docs[0]
                    )
                    if progress_callback:
                        progress_callback((i + 1) / total_chunks * 100)
            
            else:
                messagebox.showinfo("Success", f"새 문서 '{filename}'을(를) 추가합니다.")
                chunks = self.chunk_text(text, self.chunk_size)
                total_chunks = len(chunks)
                ids = [str(uuid.uuid4()) for _ in chunks]
                
                for i, chunk in enumerate(chunks):
                    collection.add(
                        documents=[chunk],
                        metadatas=[{"source": filename}],
                        ids=[ids[i]]
                    )
                    if progress_callback:
                        progress_callback((i + 1) / total_chunks * 100)       

            return total_chunks
        except Exception as e:
            error_message = f"stor_in_chroma 오류 발생: {e}"
            logger.debug(f"{error_message}")
            return []    

   
    def get_list_collections(self):
        return [col.name for col in self.client.list_collections()]
    
    def split_embed_docs_store(self, text, file_name, collection_name):
        try:
            logger.info("Starting document processing")
            chunk_size = int(self.chunk_size) if self.chunk_size is not None else 1000
            chunk_overlap = int(self.chunk_overlap) if self.chunk_overlap is not None else 200
            text_splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=chunk_overlap)
            
            chunks = text_splitter.split_documents(documents=text)
            logger.info(f"Created {len(chunks)} chunks")

            if not chunks:
                logger.warning("No chunks created from the document")
                return 0

            logger.info("Initializing embedding model")
            embeddings = self.embeddings
            logger.info("Creating Chroma vector store")
            vectordb = Chroma.from_documents(
                documents=chunks,
                embedding=embeddings,
                persist_directory=f"{self.persist_directory}/{collection_name}",
                client=self.client,
                collection_name=collection_name,
                client_settings={"anonymized_telemetry":False}
            )
            
            logger.info("Persisting vector store")
                        
            logger.info("Verifying storage")
            count = self.verify_storage(collection_name)
            
            return count
        except Exception as e:
            logger.error(f"Error in split_embed_docs_store: {str(e)}", exc_info=True)
            raise

    
        
        
    def split_keywords(self, text):
        kmat = Kkma()
        
        # 형태소 분석
        keywords = kmat.nouns(text)
        
        # 명사, 형용사, 동사만 추출
        #keywords = [word for word, pos in morphs if pos in ['Noun', 'Adjective', 'Verb']]
        
        return keywords

    # 사용 예시:
    # initial_results = [
    #     {'id': '1', 'content': 'This is the first document', 'score': 0.9, 'metadata': {'author': 'John', 'date': '2023-01-01'}},
    #     {'id': '2', 'content': 'This is the second document', 'score': 0.8, 'metadata': {'author': 'Jane', 'date': '2023-01-02'}},
    #     # ...
    # ]
    # query = "document search"
    # reranked_results = rerank_with_advanced_model(query, initial_results, top_k=5)

    #similarity_threshold 값이 높을 수록 정확도가 높아짐
    def search_collection(self, collection_name, query, n_results, source_name=None, similarity_threshold=0.7):
        """
        컬렉션을 검색하는 함수
        
        Args:
            collection_name: 검색할 컬렉션 이름
            query: 검색 쿼리
            n_results: 반환할 결과 수
            source_name: 검색할 source 이름 (선택적)
            similarity_threshold: 유사도 임계값 (기본값: 0.7)
        Returns:
            검색 결과 리스트
        """
        try:
            vectordb = self.get_or_create_collection(collection_name)
            if not vectordb:
                return []

            n_results = int(n_results) if n_results is not None else 5
            querytmp = QuoExt.extract_and_join(query)
            #print(f"query ===: {querytmp}")
            
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)
                # source_name 유무와 관계없이 동일한 메서드 사용
                results = vectordb.similarity_search_with_relevance_scores(
                    querytmp,
                    k=n_results,
                    filter={"source": source_name} if source_name else None
                )
            
            # 점수 정규화 및 결과 처리
            if results:
                scores = np.array([score for _, score in results]).reshape(-1, 1)
                scaler = MinMaxScaler()
                normalized_scores = scaler.fit_transform(scores).flatten()
                
                filtered_results = []
                for (doc, _), score in zip(results, normalized_scores):
                    if score >= similarity_threshold:
                        filtered_results.append({
                            'page_content': doc.page_content,
                            'metadata': doc.metadata,
                            'score': float(score)
                        })
                
                return filtered_results
                
        except Exception as e:
            logger.error(f"Error in search_collection: {str(e)}")
            return []
       
    
        
    def get_all_documents_source(self, collection_name, source_search):
        try:
            collection = self.client.get_collection(collection_name)
            all_docs = collection.get()
            
            if source_search is None or len(source_search) == 0:
                # source_search가 비어있으면 모든 결과 반환  100개로 제한
                all_sources = list(dict.fromkeys(metadata.get('source', 'Unknown') for metadata in all_docs['metadatas']))[:100]  # 최대 100개로 제한
                filtered_sources = []
                for s in all_sources:
                    filtered_sources.append(s)
                # logger.debug(filtered_sources)   
                return filtered_sources
            else:
                all_sources = list(dict.fromkeys(metadata.get('source', 'Unknown') for metadata in all_docs['metadatas']))
                filtered_sources = []
                for s in all_sources:
                    if source_search.lower() in s.lower():
                        if not any(fs.lower() == s.lower() for fs in filtered_sources):
                            filtered_sources.append(s)       
                return filtered_sources
        except Exception as e:
            error_message = f"get_all_documents_source 오류 발생: {e}"
            logger.debug(f"{error_message}")
            return []            
    
    def get_documents_by_source(self, collection_name, sources):
        try:
            collection = self.client.get_collection(collection_name)
            
            #logger.debug(f"Searching for documents with sources: {sources}")
            
            if isinstance(sources, str):
                sources = [sources]  # 단일 문자열을 리스트로 변환
            
            all_results = []
            
            for source in sources:
                results = collection.get(
                    where={"source": source}
                )
        
                for doc, metadata in zip(results['documents'], results['metadatas']):
                    all_results.append(Document(page_content=doc, metadata=metadata))
            
            #logger.debug(f"Found {len(all_results)} documents in total")
        
            return all_results
        except Exception as e:
            error_message = f"get_documets_by_source 오류 발생: {e}"
            logger.debug(f"{error_message}")
            return []        
        
    
    def get_ids_by_source(self, collection_name, source):
        try:
            collection = self.client.get_collection(collection_name)
            
            #logger.debug(f"Searching for document IDs with source '{source}'")
            
            results = collection.query(
                query_texts=[""],  # 빈 쿼리 텍스트
                where={"source": source},
                include=["ids"]
            )

            matching_ids = results['ids'][0]  # 'ids'는 리스트의 리스트 형태로 반환됩니다
            
            if matching_ids:
                logger.debug(f"Found {len(matching_ids)} document(s) with source '{source}'")
            else:
                logger.debug(f"No documents with source '{source}' found in the collection")
            
            return matching_ids
        except Exception as e:
            error_message = f"get_ids_by_source 오류 발생: {e}"
            logger.debug(f"{error_message}")
            return []
        
        
    def delete_source(self, collection_name: str, sources: Union[str, List[str], 'FileObject', List['FileObject']]) -> dict:
        """
        Delete documents with specific sources from a ChromaDB collection.
        
        Args:
            collection_name (str): The name of the collection.
            sources (Union[str, List[str], FileObject, List[FileObject]]): The source(s) to delete. 
                Can be a string, a file object, or a list of strings or file objects.
        
        Returns:
            dict: A dictionary containing the results of the deletion operation.
        """
        try:
            collection = self.client.get_collection(collection_name)
            
            if not isinstance(sources, list):
                sources = [sources]
            
            results = {"successful": [], "failed": []}
            
            for source in sources:
                try:
                    if not isinstance(source, str):
                        source = os.path.basename(source.name)
                    
                    delete_result = collection.delete(where={"source": str(source)})
                    
                    if delete_result:
                        logger.info(f"Successfully deleted documents with source '{source}' from collection '{collection_name}'")
                        results["successful"].append(source)
                    else:
                        logger.warning(f"No documents with source '{source}' found in collection '{collection_name}'")
                        results["failed"].append(source)
                except Exception as inner_e:
                    logger.error(f"Error deleting source '{source}' from collection '{collection_name}': {str(inner_e)}")
                    results["failed"].append(source)
            
            return results
        except Exception as e:
            logger.error(f"Error accessing collection '{collection_name}': {str(e)}")
            return {"successful": [], "failed": sources}
    
    def check_source_exists(self, collection_name, source):
        try:
            collection = self.client.get_collection(collection_name)
            if not isinstance(source, str): #파일 객체인 경우
                source = source.name   
            source = os.path.basename(source)   
           
            try:
                results = collection.get(
                    where={"source": str(source)},
                    include=["metadatas"]
                )                
                
                # metadatas가 비어있지 않으면 문서가 존재한다고 판단
                exists = len(results['metadatas']) > 0
                
                if exists:                    
                    logger.debug(f"Documents with source '{source}' exist in the collection")
                else:
                    logger.debug(f"No documents with source '{source}' found in the collection")
                            
                return exists
            except Exception as e:
                logger.debug(f"Error querying the database: {str(e)}")
                return False
        except Exception as e:
            error_message = f"check_source_exits 오류 발생: {e}"
            logger.debug(f"{error_message}")
            return False
    
    # 사용 예시
    # summary = summarize_documents_from_source("my_collection", "example_source", "gpt-3.5-turbo")

    def summarize_documents_from_source(self, collection_name, sources, llm_name, llm_model):
        def generator():
            try:
                client = self.client
                collection = client.get_collection(collection_name)
                all_documents = []
                
                total_steps = len(sources) + 3  # 소스 처리 + LLM 초기화 + 텍스트 분할 + 최종 요약
                current_step = 0

                for source in sources:
                    results = collection.get(
                        where={"source": str(source)},
                        include=["documents", "metadatas"]
                    )
                    all_documents.extend([Document(page_content=doc) for doc in results['documents']])
                    current_step += 1
                    yield {'type': 'progress', 'value': (current_step / total_steps) * 100}

                # LLM 초기화
                if llm_name == "Ollama":
                    from langchain_community.llms import Ollama
                    llm = Ollama(model=llm_model)
                elif llm_name == "Openai":
                    from langchain_community.chat_models import ChatOpenAI
                    llm = ChatOpenAI(model_name=llm_model, max_tokens=1024)
                elif llm_name == "Groq":
                    from langchain_groq import ChatGroq
                    llm = ChatGroq(model_name=llm_model, max_tokens=1024)
                else:
                    raise ValueError(f"Unsupported LLM: {llm_name}")
                
                current_step += 1
                yield {'type': 'progress', 'value': (current_step / total_steps) * 100}

                # 텍스트 분할
                text_splitter = TokenTextSplitter(chunk_size=1000, chunk_overlap=100)
                splits = text_splitter.split_documents(all_documents)
                
                current_step += 1
                yield {'type': 'progress', 'value': (current_step / total_steps) * 100}

                # 요약 프롬프트
                summarize_prompt = PromptTemplate(
                    template="다음 텍스트를 요약해주세요. 주요 포인트만 추출하여 간단명료하게 작성하세요:\n\n{text}\n\n요약:",
                    input_variables=["text"]
                )

                # 요약 체인 생성
                summarize_chain = load_summarize_chain(llm, chain_type="stuff", prompt=summarize_prompt)

                # 각 청크 요약 및 결과 저장
                summaries = []
                for i, chunk in enumerate(splits):
                    retry_count = 0
                    max_retries = 5
                    while retry_count < max_retries:
                        try:
                            summary = summarize_chain.invoke({"input_documents": [chunk]})
                            summaries.append(summary['output_text'])
                            yield {'type': 'progress', 'value': ((current_step + (i+1)/len(splits)) / total_steps) * 100}
                            break  # 성공적으로 요약되면 while 루프 종료
                        except groq.RateLimitError as e:
                            retry_count += 1
                            error_message = str(e)
                            yield {'type': 'info', 'value': f"Rate limit 오류 발생 (시도 {retry_count}/{max_retries}): {error_message}"}
                            
                            wait_time = self.parse_rate_limit_wait_time(error_message)
                            yield {'type': 'info', 'value': f"Rate limit 도달. {wait_time:.2f}초 대기 후 재시도합니다."}
                            time.sleep(wait_time + 0.5)  # 여유를 두고 조금 더 기다림
                            
                            if retry_count == max_retries:
                                yield {'type': 'error', 'value': f"최대 재시도 횟수 도달. 청크 {i+1} 요약 실패."}
                        except Exception as e:
                            error_msg = f"청크 {i+1} 요약 중 예기치 않은 오류 발생: {str(e)}\n{traceback.format_exc()}"
                            yield {'type': 'error', 'value': error_msg}
                            break  # 다른 종류의 오류 발생 시 다음 청크로 넘어감

                # 최종 요약 생성 (로컬에서 통합)
                final_summary = "\n\n".join(summaries)
            
                # 결과가 너무 길 경우 추가 요약
                if len(final_summary.split()) > 10240:  # 단어 수가 10240을 넘으면 추가 요약
                    try:
                        final_summary_doc = Document(page_content=final_summary)
                        final_summary_chain = load_summarize_chain(llm, chain_type="stuff", prompt=summarize_prompt)
                        final_summary = final_summary_chain.invoke({"input_documents": [final_summary_doc]})['output_text']
                    except Exception as e:
                        error_msg = f"최종 요약 생성 중 오류 발생: {str(e)}\n{traceback.format_exc()}"
                        yield {'type': 'error', 'value': error_msg}
                        # 오류 발생 시 기존의 요약을 그대로 사용

                yield {'type': 'summary', 'value': final_summary}
                #print(f"요약결과  : {final_summary}")
            
            except Exception as e:
                error_msg = f"요약 프로세스 중 예기치 않은 오류 발생: {str(e)}\n{traceback.format_exc()}"
                yield {'type': 'error', 'value': error_msg}
        return generator()
    
    def parse_rate_limit_wait_time(self, error_message):
        # 정규 표현식을 사용하여 대기 시간 추출
        match = re.search(r'Please try again in (\d+m)?(\d+(\.\d+)?)s', error_message)
        if match:
            minutes = int(match.group(1)[:-1]) if match.group(1) else 0
            seconds = float(match.group(2))
            return minutes * 60 + seconds
        else:
            # 대기 시간을 추출할 수 없는 경우 기본값 반환
            return 60  # 기본 1분 대기

    def close(self):
        # 필요한 정리 작업을 수행
        pass